import csv
import io
import json
import os
import re
import secrets
import sqlite3
from contextlib import contextmanager
import urllib.error
import urllib.request
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from flask import Flask, Response, abort, jsonify, render_template, request, session
from werkzeug.utils import secure_filename

try:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
except ImportError:
    Document = None
    PackageNotFoundError = None
    PdfReader = None
    PdfReadError = None


BASE_DIR = Path(__file__).resolve().parent
DATABASE = Path(os.environ.get("RESUME_ANALYZER_DB", BASE_DIR / "resume_analyzer.db"))
ALLOWED_EXTENSIONS = {".pdf", ".docx"}
MAX_UPLOAD_BYTES = 8 * 1024 * 1024
SKILL_ALIASES = {
    "python": ("python",),
    "java": ("java",),
    "javascript": ("javascript", "js"),
    "typescript": ("typescript",),
    "sql": ("sql", "sqlite", "mysql", "postgresql", "postgres"),
    "flask": ("flask",),
    "django": ("django",),
    "react": ("react",),
    "node.js": ("node.js", "nodejs"),
    "html/css": ("html", "css"),
    "rest apis": ("rest api", "restful api", "rest apis"),
    "git": ("git", "github"),
    "machine learning": ("machine learning", "scikit-learn", "tensorflow", "pytorch"),
    "data analysis": ("data analysis", "data analytics", "pandas", "power bi", "tableau"),
    "aws": ("aws", "amazon web services"),
    "docker": ("docker",),
    "communication": ("communication", "presented", "presentation"),
    "teamwork": ("teamwork", "collaborated", "collaboration"),
    "problem solving": ("problem solving", "analytical thinking"),
    "leadership": ("leadership", "led a team", "team lead"),
    "time management": ("time management", "prioritization"),
}
SOFT_SKILLS = {"communication", "teamwork", "problem solving", "leadership", "time management"}
LEARNING = {
    "python": {"course": "Python for Everybody", "provider": "Coursera", "project": "Build a command-line study planner with file storage."},
    "java": {"course": "Java Programming and Software Engineering Fundamentals", "provider": "Coursera", "project": "Create a Java REST service with input validation and tests."},
    "javascript": {"course": "JavaScript Algorithms and Data Structures", "provider": "freeCodeCamp", "project": "Build an accessible interactive application using browser APIs."},
    "typescript": {"course": "TypeScript Handbook", "provider": "Microsoft Learn", "project": "Convert a JavaScript project to strict TypeScript."},
    "sql": {"course": "SQLBolt interactive lessons", "provider": "SQLBolt", "project": "Design a relational database and write reporting queries."},
    "flask": {"course": "Flask Mega-Tutorial", "provider": "Miguel Grinberg", "project": "Build a small Flask app with authentication and database migrations."},
    "django": {"course": "Django for Everybody", "provider": "Coursera", "project": "Create a Django directory app with search and pagination."},
    "react": {"course": "React Learn", "provider": "React.dev", "project": "Build a component-based dashboard with accessible forms."},
    "node.js": {"course": "Introduction to Node.js", "provider": "Node.js", "project": "Build a JSON API with validation and automated tests."},
    "html/css": {"course": "Responsive Web Design", "provider": "freeCodeCamp", "project": "Create a responsive, keyboard-accessible portfolio."},
    "rest apis": {"course": "API design learning path", "provider": "Postman Academy", "project": "Document and test a versioned REST API."},
    "git": {"course": "Introduction to GitHub", "provider": "GitHub Skills", "project": "Publish a project with a clean history and collaborative pull requests."},
    "machine learning": {"course": "Machine Learning Specialization", "provider": "DeepLearning.AI / Coursera", "project": "Train and evaluate a documented classification model."},
    "data analysis": {"course": "Data Analysis with Python", "provider": "freeCodeCamp", "project": "Analyze an open dataset and publish a reproducible notebook."},
    "aws": {"course": "AWS Educate learning paths", "provider": "AWS Educate", "project": "Deploy a small application using a free-tier architecture."},
    "docker": {"course": "Docker Get Started", "provider": "Docker Docs", "project": "Containerize an app and document its local setup."},
    "communication": {"course": "Career communication resources", "provider": "LinkedIn Learning", "project": "Present a project demo and write a one-page technical brief."},
    "teamwork": {"course": "Teamwork skills resources", "provider": "Coursera", "project": "Contribute to a team project with documented responsibilities."},
    "problem solving": {"course": "Problem Solving Through Programming", "provider": "Coursera", "project": "Solve and explain a set of algorithmic problems."},
    "leadership": {"course": "Leadership skills resources", "provider": "OpenLearn", "project": "Coordinate a small student project and document outcomes."},
    "time management": {"course": "Personal productivity resources", "provider": "Microsoft Learn", "project": "Plan a short project using milestones and a retrospective."},
}

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY") or secrets.token_hex(32)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES


@contextmanager
def connect_db():
    connection = sqlite3.connect(DATABASE)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        with connection:
            yield connection
    finally:
        connection.close()


def init_db():
    DATABASE.parent.mkdir(parents=True, exist_ok=True)
    with connect_db() as db:
        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS candidates (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                email TEXT NOT NULL DEFAULT '',
                filename TEXT NOT NULL,
                resume_text TEXT NOT NULL,
                skills_json TEXT NOT NULL,
                unverified_json TEXT NOT NULL DEFAULT '[]',
                education TEXT NOT NULL DEFAULT '',
                experience TEXT NOT NULL DEFAULT '',
                projects TEXT NOT NULL DEFAULT '',
                strengths TEXT NOT NULL DEFAULT '',
                summary TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS jobs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                department TEXT NOT NULL DEFAULT '',
                description TEXT NOT NULL DEFAULT '',
                required_json TEXT NOT NULL,
                preferred_json TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS shortlists (
                job_id INTEGER NOT NULL REFERENCES jobs(id) ON DELETE CASCADE,
                candidate_id INTEGER NOT NULL REFERENCES candidates(id) ON DELETE CASCADE,
                created_at TEXT NOT NULL,
                PRIMARY KEY (job_id, candidate_id)
            );
            """
        )
        columns = {row["name"] for row in db.execute("PRAGMA table_info(candidates)")}
        if "unverified_json" not in columns:
            db.execute("ALTER TABLE candidates ADD COLUMN unverified_json TEXT NOT NULL DEFAULT '[]'")
        if db.execute("SELECT COUNT(*) FROM jobs").fetchone()[0] == 0:
            now = timestamp()
            demo_jobs = [
                ("Junior Python Developer", "Engineering", "Build and test web APIs with Python, Flask, SQL and Git. Clear communication and teamwork are valued.", ["python", "flask", "sql", "git"], ["rest apis", "communication", "teamwork"]),
                ("Data Analyst Intern", "Analytics", "Use Python, SQL, and data analysis to communicate insights. Strong problem solving and presentation skills preferred.", ["python", "sql", "data analysis"], ["communication", "problem solving", "teamwork"]),
                ("Frontend Developer Intern", "Product", "Create responsive web experiences with JavaScript, React, HTML/CSS and Git.", ["javascript", "react", "html/css", "git"], ["communication", "teamwork"]),
            ]
            db.executemany(
                "INSERT INTO jobs (title, department, description, required_json, preferred_json, created_at) VALUES (?, ?, ?, ?, ?, ?)",
                [(title, dept, desc, json.dumps(required), json.dumps(preferred), now) for title, dept, desc, required, preferred in demo_jobs],
            )
        if db.execute("SELECT COUNT(*) FROM candidates").fetchone()[0] == 0:
            seed_candidates(db)


def timestamp():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def seed_candidates(db):
    examples = [
        ("Aarav Mehta", "aarav@example.edu", "Sample resume", "Computer Science student with experience building a Flask API in Python and SQL. Used Git for team projects. Communicated project results and collaborated with teammates. Education: B.Tech Computer Science. Projects: Student resource portal.", ["python", "flask", "sql", "git", "communication", "teamwork"], "B.Tech Computer Science", "Flask API and team software projects", "Student resource portal", "API development, collaboration", "Computer Science student with hands-on Python and Flask experience, SQL data work, and collaborative project delivery."),
        ("Diya Rao", "diya@example.edu", "Sample resume", "Information Science student. Analyzed datasets with Python, pandas and SQL; created dashboards in Tableau. Presented findings and solved data quality problems. Education: B.E. Information Science. Projects: Retail trend analysis.", ["python", "sql", "data analysis", "communication", "problem solving"], "B.E. Information Science", "Data analysis and dashboard projects", "Retail trend analysis", "Data storytelling, analytical thinking", "Information Science student with practical Python and SQL analysis experience and a focus on communicating data insights."),
        ("Kabir Shah", "kabir@example.edu", "Sample resume", "Web development student experienced with JavaScript, React, HTML and CSS. Built responsive interfaces with Git and worked in a team. Education: B.Tech Information Technology. Projects: Campus events app.", ["javascript", "react", "html/css", "git", "teamwork"], "B.Tech Information Technology", "Frontend development projects", "Campus events app", "Responsive UI, teamwork", "Information Technology student with React and JavaScript experience building responsive web interfaces."),
    ]
    for name, email, filename, text, skills, education, experience, projects, strengths, summary in examples:
        db.execute(
            """INSERT INTO candidates (name, email, filename, resume_text, skills_json, education,
               experience, projects, strengths, summary, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (name, email, filename, text, json.dumps(skills), education, experience, projects, strengths, summary, timestamp()),
        )


def clean_text(value, limit=12000):
    return re.sub(r"\s+", " ", str(value or "")).strip()[:limit]


def clean_resume_text(value):
    return "\n".join(re.sub(r"[ \t]+", " ", line).strip() for line in value.splitlines() if line.strip())[:40000]


def extract_resume(upload):
    suffix = Path(secure_filename(upload.filename or "")).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("Upload a PDF or DOCX resume.")
    if suffix == ".pdf":
        if PdfReader is None:
            raise RuntimeError("PDF support is unavailable. Install the packages in requirements.txt.")
        reader = PdfReader(upload.stream)
        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported.")
        text = "\n".join(page.extract_text() or "" for page in reader.pages[:30])
    else:
        if Document is None:
            raise RuntimeError("DOCX support is unavailable. Install the packages in requirements.txt.")
        document = Document(upload.stream)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            text += "\n" + "\n".join(" ".join(cell.text for cell in row.cells) for row in table.rows)
    text = clean_resume_text(text)
    if len(text) < 40:
        raise ValueError("We couldn't read enough text from this resume. Try a text-based PDF or DOCX.")
    return text


def contains_skill(text, aliases):
    for alias in aliases:
        pattern = r"(?<![a-z0-9])" + re.escape(alias).replace(r"\ ", r"\s+") + r"(?![a-z0-9])"
        if re.search(pattern, text, re.IGNORECASE):
            return True
    return False


def analyze_resume(text):
    normalized = text.lower()
    found = [skill for skill, aliases in SKILL_ALIASES.items() if contains_skill(normalized, aliases)]
    unverified = []
    weak_evidence = re.compile(r"\b(familiar with|exposure to|basic knowledge of|beginner in|currently learning|learning)\b", re.IGNORECASE)
    strong_evidence = re.compile(r"\b(used|built|implemented|developed|created|designed|deployed|proficient|expert|experience(?:d)? with|worked with|applied)\b", re.IGNORECASE)
    for skill in found:
        weak_mention = False
        strong_mention = False
        for alias in SKILL_ALIASES[skill]:
            matches = re.finditer(r"(?<![a-z0-9])" + re.escape(alias).replace(r"\ ", r"\s+") + r"(?![a-z0-9])", normalized, re.IGNORECASE)
            for match in matches:
                boundaries = [normalized.rfind(mark, 0, match.start()) for mark in ".!?;\n"]
                sentence_start = max(boundaries) + 1
                before = normalized[sentence_start:match.start()]
                weak_mention = weak_mention or bool(weak_evidence.search(before))
                strong_mention = strong_mention or bool(strong_evidence.search(before))
        if weak_mention and not strong_mention:
            unverified.append(skill)
    sections = {}
    for label, headings in {
        "education": ("education", "academic background"),
        "experience": ("experience", "employment", "work history", "internship"),
        "projects": ("projects", "project experience"),
    }.items():
        lines = [line.strip(" •-\t") for line in text.splitlines() if line.strip()]
        selections = []
        collecting = False
        for line in lines:
            lowered = line.lower().strip(" :")
            if any(lowered.startswith(heading) for heading in headings):
                collecting = True
                remainder = line.split(":", 1)[-1].strip()
                if remainder and remainder.lower() != lowered:
                    selections.append(remainder)
                continue
            if collecting and re.match(r"^[A-Z][A-Za-z /&-]{2,30}:?$", line.strip()):
                break
            if collecting:
                selections.append(line)
            if len(selections) >= 3:
                break
        sections[label] = clean_text(" · ".join(selections), 280)
    strengths = []
    for skill in found:
        if skill in SOFT_SKILLS:
            continue
        strengths.append(skill.title())
    if not strengths:
        strengths = [skill.title() for skill in found[:4]]
    if not strengths:
        strengths = ["Relevant experience may need manual review"]
    summary_parts = []
    if sections["education"]:
        summary_parts.append("Education: " + sections["education"] + ".")
    if sections["experience"]:
        summary_parts.append("Experience: " + sections["experience"] + ".")
    if sections["projects"]:
        summary_parts.append("Projects: " + sections["projects"] + ".")
    if found:
        summary_parts.append("Skills identified: " + ", ".join(found) + ".")
    summary = " ".join(summary_parts) or "Resume text was extracted, but structured education, experience, and project details were not clearly labeled. Review the original resume."
    return {
        "skills": found,
        "unverified": unverified,
        "education": sections["education"] or "Not clearly identified",
        "experience": sections["experience"] or "Not clearly identified",
        "projects": sections["projects"] or "Not clearly identified",
        "strengths": ", ".join(strengths),
        "summary": summary,
    }


def job_dict(row):
    return {
        "id": row["id"], "title": row["title"], "department": row["department"],
        "description": row["description"], "required": json.loads(row["required_json"]),
        "preferred": json.loads(row["preferred_json"]), "created_at": row["created_at"],
    }


def candidate_dict(row):
    return {
        "id": row["id"], "name": row["name"], "email": row["email"], "filename": row["filename"],
        "skills": json.loads(row["skills_json"]), "unverified": json.loads(row["unverified_json"]),
        "technical_skills": [skill for skill in json.loads(row["skills_json"]) if skill not in SOFT_SKILLS],
        "soft_skills": [skill for skill in json.loads(row["skills_json"]) if skill in SOFT_SKILLS],
        "education": row["education"],
        "experience": row["experience"], "projects": row["projects"], "strengths": row["strengths"],
        "summary": row["summary"], "created_at": row["created_at"],
    }


def score_candidate(candidate, job):
    unverified = {skill.lower() for skill in candidate.get("unverified", [])}
    skills = {skill.lower() for skill in candidate["skills"] if skill.lower() not in unverified}
    required = job["required"]
    preferred = job["preferred"]
    required_keys = [canonical_skill(skill) for skill in required]
    preferred_keys = [canonical_skill(skill) for skill in preferred]
    requirement_results = []
    for skill in required + preferred:
        key = canonical_skill(skill)
        known = key in skills
        required_skill = skill in required
        requirement_results.append({
            "skill": skill, "type": "Essential" if required_skill else "Preferred",
            "status": "matched" if known else ("unverified" if key in unverified else "missing"),
        })
    required_rate = sum(skill in skills for skill in required_keys) / len(required) if required else 1
    preferred_rate = sum(skill in skills for skill in preferred_keys) / len(preferred) if preferred else 1
    score = int(((required_rate * 75 + preferred_rate * 25) if required and preferred else (required_rate * 100 if required else preferred_rate * 100)) + 0.5)
    matched = [item["skill"] for item in requirement_results if item["status"] == "matched"]
    missing = [item["skill"] for item in requirement_results if item["status"] in {"missing", "unverified"}]
    required_matched = [skill for skill, key in zip(required, required_keys) if key in skills]
    required_missing = [skill for skill, key in zip(required, required_keys) if key not in skills and key not in unverified]
    required_unverified = [skill for skill, key in zip(required, required_keys) if key in unverified]
    verified_note = f" {sum(item['status'] == 'unverified' for item in requirement_results)} mention(s) need verification." if any(item["status"] == "unverified" for item in requirement_results) else ""
    explanation = f"{len(required_matched)}/{len(required)} essential and {sum(skill in skills for skill in preferred_keys)}/{len(preferred)} preferred skills found. Essential skills count for 75% and preferred skills for 25% of the score.{verified_note}"
    return {
        "score": score, "matched": matched, "missing": missing, "requirements": requirement_results,
        "required_matched": required_matched, "required_missing": required_missing,
        "required_unverified": required_unverified,
        "unverified": [item["skill"] for item in requirement_results if item["status"] == "unverified"],
        "explanation": explanation,
    }


def get_optional_ai_summary(candidate, job):
    api_key = os.environ.get("OPENAI_API_KEY")
    if not api_key:
        return candidate["summary"]
    endpoint = os.environ.get("OPENAI_BASE_URL", "https://api.openai.com/v1/chat/completions").rstrip("/")
    payload = {
        "model": os.environ.get("OPENAI_MODEL", "gpt-4o-mini"),
        "messages": [
            {"role": "system", "content": "Write a concise, evidence-based resume summary for a college placement officer. Only state facts present in the resume. Never infer demographic traits. Note uncertainty rather than inventing details."},
            {"role": "user", "content": f"Job: {job['title'] if job else 'General placement'}\nResume text:\n{candidate.get('_resume_text', '')[:10000]}"},
        ],
        "temperature": 0.2,
    }
    req = urllib.request.Request(
        endpoint + "/chat/completions", data=json.dumps(payload).encode(),
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}, method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=15) as response:
            result = json.loads(response.read().decode())
        content = result["choices"][0]["message"]["content"].strip()
        return content[:2500] or candidate["summary"]
    except (urllib.error.URLError, TimeoutError, ValueError, KeyError, IndexError, TypeError, AttributeError) as error:
        app.logger.warning("Optional AI summary unavailable: %s", error)
        return candidate["summary"] + " (Local demo summary; AI provider was unavailable.)"


def shortlist_ids(job_id):
    with connect_db() as db:
        return {row["candidate_id"] for row in db.execute("SELECT candidate_id FROM shortlists WHERE job_id = ?", (job_id,))}


def candidate_for_job(row, job, shortlisted=None):
    result = candidate_dict(row)
    result["match"] = score_candidate(result, job)
    result["shortlisted"] = result["id"] in shortlisted if shortlisted is not None else False
    return result


def csrf_required():
    expected = session.get("csrf_token")
    supplied = request.headers.get("X-CSRF-Token", "")
    if not expected or not secrets.compare_digest(expected, supplied):
        return jsonify({"error": "Your session token expired. Refresh the page and try again."}), 403
    return None


@app.before_request
def ensure_database():
    if not getattr(app, "_database_ready", False):
        init_db()
        app._database_ready = True


@app.before_request
def protect_api():
    if request.path.startswith("/api/") and request.method in {"POST", "DELETE"}:
        return csrf_required()
    return None


@app.errorhandler(413)
def file_too_large(_error):
    return jsonify({"error": "File is too large. Maximum upload size is 8 MB."}), 413


@app.get("/")
def home():
    session.setdefault("csrf_token", secrets.token_urlsafe(32))
    return render_template("index.html", csrf_token=session["csrf_token"])


@app.get("/api/dashboard")
def dashboard():
    with connect_db() as db:
        candidates = db.execute("SELECT * FROM candidates ORDER BY created_at DESC").fetchall()
        jobs = [job_dict(row) for row in db.execute("SELECT * FROM jobs ORDER BY created_at DESC")]
        shortlist_count = db.execute("SELECT COUNT(*) FROM shortlists").fetchone()[0]
    top = []
    for job in jobs:
        for row in candidates:
            candidate = candidate_for_job(row, job)
            top.append({"candidate": candidate, "job": {"id": job["id"], "title": job["title"]}})
    top.sort(key=lambda item: item["candidate"]["match"]["score"], reverse=True)
    return jsonify({
        "candidate_count": len(candidates), "job_count": len(jobs), "shortlist_count": shortlist_count,
        "average_match": round(sum(item["candidate"]["match"]["score"] for item in top) / len(top)) if top else 0,
        "jobs": jobs, "top_matches": top[:5],
    })


@app.get("/api/jobs")
def list_jobs():
    with connect_db() as db:
        rows = db.execute("SELECT * FROM jobs ORDER BY created_at DESC").fetchall()
    return jsonify([job_dict(row) for row in rows])


@app.post("/api/jobs")
def create_job():
    data = request.get_json(silent=True) or {}
    title = clean_text(data.get("title"), 120)
    department = clean_text(data.get("department"), 80)
    description = clean_text(data.get("description"), 3000)
    required = parse_skill_list(data.get("required"))
    preferred = parse_skill_list(data.get("preferred"))
    inferred_required, inferred_preferred, inferred_other = extract_job_requirements(description)
    if not required:
        required.extend(inferred_required + inferred_other)
    else:
        required.extend(skill for skill in inferred_required if skill not in required)
    preferred.extend(skill for skill in inferred_preferred if skill not in preferred and skill not in required)
    if not title or not required:
        return jsonify({"error": "A role title and at least one essential skill are required."}), 400
    with connect_db() as db:
        cursor = db.execute(
            "INSERT INTO jobs (title, department, description, required_json, preferred_json, created_at) VALUES (?, ?, ?, ?, ?, ?)",
            (title, department, description, json.dumps(required), json.dumps(preferred), timestamp()),
        )
        row = db.execute("SELECT * FROM jobs WHERE id = ?", (cursor.lastrowid,)).fetchone()
    return jsonify(job_dict(row)), 201


def parse_skill_list(value):
    values = value if isinstance(value, list) else str(value or "").split(",")
    result = []
    for item in values:
        skill = clean_text(item, 60).lower()
        if skill and skill not in result:
            result.append(skill)
    return result[:30]


def canonical_skill(value):
    normalized = clean_text(value, 80).lower()
    for skill, aliases in SKILL_ALIASES.items():
        if normalized == skill or normalized in aliases:
            return skill
    return normalized


def extract_job_requirements(description):
    text = description.lower()
    essential, preferred, unspecified = [], [], []
    essential_cue = re.compile(r"\b(required|essential|must[- ]have|must possess|qualification)\b")
    preferred_cue = re.compile(r"\b(preferred|nice to have|bonus|desirable|a plus)\b")
    for skill, aliases in SKILL_ALIASES.items():
        match = None
        for alias in aliases:
            match = re.search(r"(?<![a-z0-9])" + re.escape(alias).replace(r"\ ", r"\s+") + r"(?![a-z0-9])", text)
            if match:
                break
        if not match:
            continue
        sentence = next((item for item in re.finditer(r"[^.!?;]+", text) if item.start() <= match.start() < item.end()), None)
        context = sentence.group(0) if sentence else text[max(0, match.start() - 70):match.end() + 55]
        relative_position = match.start() - (sentence.start() if sentence else max(0, match.start() - 70))
        cues = [(abs(cue.start() - relative_position), category) for category, pattern in (("essential", essential_cue), ("preferred", preferred_cue)) for cue in pattern.finditer(context)]
        if cues and min(cues)[1] == "preferred":
            preferred.append(skill)
        elif cues and min(cues)[1] == "essential":
            essential.append(skill)
        else:
            unspecified.append(skill)
    return essential, preferred, unspecified


@app.get("/api/candidates")
def list_candidates():
    search = clean_text(request.args.get("search"), 120).lower()
    job_id = request.args.get("job_id", type=int)
    with connect_db() as db:
        rows = db.execute("SELECT * FROM candidates ORDER BY created_at DESC").fetchall()
        jobs = [job_dict(row) for row in db.execute("SELECT * FROM jobs")]
    job_map = {job["id"]: job for job in jobs}
    job = job_map.get(job_id) if job_id else None
    shortlisted = shortlist_ids(job_id) if job_id else set()
    results = []
    for row in rows:
        candidate = candidate_for_job(row, job, shortlisted) if job else candidate_dict(row)
        if search and search not in " ".join([candidate["name"], candidate["email"], candidate["summary"], *candidate["skills"]]).lower():
            continue
        results.append(candidate)
    if job:
        results.sort(key=lambda candidate: candidate["match"]["score"], reverse=True)
    return jsonify(results)


@app.post("/api/candidates")
def upload_candidate():
    upload = request.files.get("resume")
    if upload is None or not upload.filename:
        return jsonify({"error": "Choose a PDF or DOCX resume to upload."}), 400
    name = clean_text(request.form.get("name"), 120)
    email = clean_text(request.form.get("email"), 160)
    if not name:
        return jsonify({"error": "Student name is required."}), 400
    if email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        return jsonify({"error": "Enter a valid email address."}), 400
    try:
        text = extract_resume(upload)
    except (ValueError, RuntimeError, OSError, zipfile.BadZipFile) as error:
        return jsonify({"error": str(error)}), 400
    except tuple(error_type for error_type in (PdfReadError, PackageNotFoundError) if error_type is not None) as error:
        return jsonify({"error": f"This file could not be read as a valid resume: {error}"}), 400
    analysis = analyze_resume(text)
    with connect_db() as db:
        cursor = db.execute(
            """INSERT INTO candidates (name, email, filename, resume_text, skills_json, unverified_json,
               education, experience, projects, strengths, summary, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (name, email, secure_filename(upload.filename)[:180], text, json.dumps(analysis["skills"]),
             json.dumps(analysis["unverified"]), analysis["education"], analysis["experience"], analysis["projects"], analysis["strengths"],
             analysis["summary"], timestamp()),
        )
        row = db.execute("SELECT * FROM candidates WHERE id = ?", (cursor.lastrowid,)).fetchone()
    candidate = candidate_dict(row)
    return jsonify(candidate), 201


@app.get("/api/candidates/<int:candidate_id>")
def get_candidate(candidate_id):
    job_id = request.args.get("job_id", type=int)
    with connect_db() as db:
        row = db.execute("SELECT * FROM candidates WHERE id = ?", (candidate_id,)).fetchone()
        job_row = db.execute("SELECT * FROM jobs WHERE id = ?", (job_id,)).fetchone() if job_id else None
    if row is None:
        abort(404)
    candidate = candidate_dict(row)
    candidate["_resume_text"] = row["resume_text"]
    if job_row:
        job = job_dict(job_row)
        candidate["match"] = score_candidate(candidate, job)
        candidate["ai_summary"] = get_optional_ai_summary(candidate, job)
        with connect_db() as db:
            all_jobs = [job_dict(item) for item in db.execute("SELECT * FROM jobs")]
        candidate["recommendations"] = recommend_jobs(candidate, all_jobs)
        candidate["learning"] = learning_for_skills(candidate["match"]["missing"])
        candidate["improvements"] = improvements(candidate, candidate["match"]["missing"])
    else:
        with connect_db() as db:
            jobs = [job_dict(item) for item in db.execute("SELECT * FROM jobs")]
        candidate["recommendations"] = recommend_jobs(candidate, jobs)
        best_missing = candidate["recommendations"][0]["missing"] if candidate["recommendations"] else []
        candidate["learning"] = learning_for_skills(best_missing)
        candidate["improvements"] = improvements(candidate, best_missing)
    candidate.pop("_resume_text", None)
    return jsonify(candidate)


def recommend_jobs(candidate, jobs):
    suggestions = []
    for job in jobs:
        match = score_candidate(candidate, job)
        suggestions.append({"id": job["id"], "title": job["title"], "department": job["department"], "score": match["score"], "matched": match["matched"], "missing": match["missing"]})
    return sorted(suggestions, key=lambda item: item["score"], reverse=True)[:5]


def learning_for_skills(skills):
    links = {
        "python": "https://www.coursera.org/specializations/python",
        "java": "https://www.coursera.org/specializations/java-programming",
        "javascript": "https://www.freecodecamp.org/learn/javascript-algorithms-and-data-structures-v8/",
        "typescript": "https://www.typescriptlang.org/docs/handbook/intro.html",
        "sql": "https://sqlbolt.com/",
        "flask": "https://blog.miguelgrinberg.com/post/the-flask-mega-tutorial-part-i-hello-world",
        "django": "https://www.coursera.org/specializations/django",
        "react": "https://react.dev/learn",
        "node.js": "https://nodejs.org/en/learn/getting-started/introduction-to-nodejs",
        "html/css": "https://www.freecodecamp.org/learn/2022/responsive-web-design/",
        "rest apis": "https://academy.postman.com/",
        "git": "https://skills.github.com/",
        "machine learning": "https://www.coursera.org/specializations/machine-learning-introduction",
        "data analysis": "https://www.freecodecamp.org/learn/data-analysis-with-python/",
        "aws": "https://aws.amazon.com/training/learn-about/educate/",
        "docker": "https://docs.docker.com/get-started/",
        "communication": "https://www.coursera.org/learn/wharton-communication-skills",
        "teamwork": "https://www.coursera.org/learn/teamwork-skills-effective-communication",
        "problem solving": "https://www.coursera.org/learn/problem-solving",
        "leadership": "https://www.open.edu/openlearn/money-business/leadership-management",
        "time management": "https://support.microsoft.com/en-us/topic/time-management-tips-for-students-7bdc24bb-64d1-4ad9-824f-4bb30a7a9cf4",
    }
    return [{"skill": skill, **LEARNING[skill], "url": links[skill]} for skill in skills if skill in LEARNING]


def improvements(candidate, missing):
    tips = []
    for skill in missing[:4]:
        tips.append(f"Build evidence for {skill}: add a small project or coursework example and describe your own contribution.")
    if candidate["education"] == "Not clearly identified":
        tips.append("Add a clearly labeled Education section with your degree, institution, and expected graduation date.")
    if candidate["experience"] == "Not clearly identified":
        tips.append("Make internships, volunteering, and relevant work easy to scan in a labeled Experience section.")
    if candidate["projects"] == "Not clearly identified":
        tips.append("Add 2–3 projects with the problem, your contribution, technologies, and measurable outcome.")
    if len(candidate["summary"]) > 0 and not candidate["email"]:
        tips.append("Add a professional contact email so placement officers can follow up.")
    return tips[:6] or ["Tailor your resume summary and project evidence to the essential skills in the selected role."]


@app.post("/api/compare")
def compare_candidates():
    data = request.get_json(silent=True) or {}
    job_id = data.get("job_id")
    ids = data.get("candidate_ids")
    if not isinstance(ids, list) or not 2 <= len(ids) <= 4 or not isinstance(job_id, int):
        return jsonify({"error": "Select 2–4 applicants and a role to compare."}), 400
    with connect_db() as db:
        job_row = db.execute("SELECT * FROM jobs WHERE id = ?", (job_id,)).fetchone()
        rows = db.execute(f"SELECT * FROM candidates WHERE id IN ({','.join('?' for _ in ids)})", ids).fetchall()
    if job_row is None or len(rows) != len(set(ids)):
        return jsonify({"error": "The selected role or applicants could not be found."}), 404
    job = job_dict(job_row)
    results = [candidate_for_job(row, job) for row in rows]
    results.sort(key=lambda candidate: candidate["match"]["score"], reverse=True)
    return jsonify({"job": job, "candidates": results})


@app.post("/api/shortlist")
def toggle_shortlist():
    data = request.get_json(silent=True) or {}
    candidate_id, job_id = data.get("candidate_id"), data.get("job_id")
    if not isinstance(candidate_id, int) or not isinstance(job_id, int):
        return jsonify({"error": "Choose an applicant and a role."}), 400
    with connect_db() as db:
        if not db.execute("SELECT 1 FROM candidates WHERE id = ?", (candidate_id,)).fetchone():
            return jsonify({"error": "Applicant not found."}), 404
        if not db.execute("SELECT 1 FROM jobs WHERE id = ?", (job_id,)).fetchone():
            return jsonify({"error": "Role not found."}), 404
        existing = db.execute("SELECT 1 FROM shortlists WHERE job_id = ? AND candidate_id = ?", (job_id, candidate_id)).fetchone()
        if existing:
            db.execute("DELETE FROM shortlists WHERE job_id = ? AND candidate_id = ?", (job_id, candidate_id))
            shortlisted = False
        else:
            db.execute("INSERT INTO shortlists (job_id, candidate_id, created_at) VALUES (?, ?, ?)", (job_id, candidate_id, timestamp()))
            shortlisted = True
    return jsonify({"shortlisted": shortlisted})


@app.get("/api/reports.csv")
def report_csv():
    job_id = request.args.get("job_id", type=int)
    if not job_id:
        return jsonify({"error": "Choose a role before exporting a report."}), 400
    with connect_db() as db:
        job_row = db.execute("SELECT * FROM jobs WHERE id = ?", (job_id,)).fetchone()
        rows = db.execute("SELECT * FROM candidates ORDER BY name").fetchall()
    if job_row is None:
        return jsonify({"error": "Role not found."}), 404
    job = job_dict(job_row)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Applicant", "Email", "Role", "Match score", "Essential matched", "Essential missing", "Unverified essential", "Preferred matched", "Shortlisted"])
    shortlisted = shortlist_ids(job_id)
    for row in rows:
        candidate = candidate_dict(row)
        match = score_candidate(candidate, job)
        writer.writerow([candidate["name"], candidate["email"], job["title"], match["score"], ", ".join(match["required_matched"]), ", ".join(match["required_missing"]), ", ".join(match["required_unverified"]), ", ".join(skill for skill in match["matched"] if skill not in match["required_matched"]), "Yes" if candidate["id"] in shortlisted else "No"])
    return Response("\ufeff" + output.getvalue(), mimetype="text/csv", headers={"Content-Disposition": f'attachment; filename="placement-report-{job_id}.csv"'})


if __name__ == "__main__":
    init_db()
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "5000")), debug=False)
